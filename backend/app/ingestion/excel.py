"""Adapters for the supplied 1C exports. Never evaluate workbook formulas."""
import hashlib
import io
import math
import re
import zipfile
from collections import defaultdict
from datetime import date, datetime, timezone
from pathlib import PurePosixPath

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from . import ImportFailure

from app.contracts import (
    Dataset, Inbound, Issue, MonthlyObservation, NormalizedDataset, OrderConstraint,
    Product, SeasonalFactor, SourceRef, StockSnapshot, StockoutInterval, Transaction, TransactionCoverage,
)

MONTHS = ['янв', 'фев', 'мар', 'апр', 'май', 'июн', 'июл', 'авг', 'сен', 'окт', 'ноя', 'дек']


def text(value):
    if value is None:
        return ''
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()


def number(value):
    if value is None or value == '':
        return None
    if isinstance(value, bool):
        return None
    try:
        result = float(str(value).replace('\u00a0', '').replace(' ', '').replace(',', '.'))
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def full_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = re.search(r'(\d{2})\.(\d{2})\.(20\d{2})', text(value))
    if match:
        try:
            return date(int(match[3]), int(match[2]), int(match[1]))
        except ValueError:
            return None
    try:
        return date.fromisoformat(text(value))
    except ValueError:
        pass
    return None


def month_date(value):
    value = text(value).lower()
    year = re.search(r'20\d{2}', value)
    if year:
        for index, prefix in enumerate(MONTHS, 1):
            if value.startswith(prefix):
                return date(int(year[0]), index, 1)
    return None


def incoming_date(value, year):
    # Comments override a header; the report year is explicitly documented.
    exact = full_date(value)
    if exact:
        return exact
    match = re.search(r'(\d{2})\.(\d{2})(?!\d)', text(value))
    if match and year:
        try:
            return date(year, int(match[2]), int(match[1]))
        except ValueError:
            pass
    return None


def unit(value):
    value = text(value).lower().rstrip('.')
    return {'шт': 'шт', 'м': 'м', 'упак': 'упак'}.get(value, value or None)


def identify(rows, supplier, filename):
    first = [text(v).lower() for v in rows[0]]
    second = [text(v).lower() for v in rows[1]] if len(rows) > 1 else []
    if first[:4] == ['дата', 'номер', 'документ', 'код']:
        return 'transactions'
    if first[:3] == ['sku', 'start', 'end']:
        return 'stockouts'
    if len(second) > 51 and second[51] == 'свободный остаток':
        if supplier != 'systeme':
            raise ImportFailure('Снимок Systeme не соответствует выбранному поставщику')
        return 'snapshot'
    if 'мин. разр. к отгр.' in first:
        if supplier != 'iek':
            raise ImportFailure('MOQ IEK не соответствует выбранному поставщику')
        return 'constraints'
    if len(first) == 5 and first[-1] == 'кратность' and first[2] == 'номенклатура.код':
        if supplier != 'systeme':
            raise ImportFailure('Кратность Systeme не соответствует выбранному поставщику')
        return 'constraints'
    if first and first[0] == 'код 1с' and any('поступление до' in h for h in first):
        if supplier != 'iek':
            raise ImportFailure('Поставки IEK не соответствуют выбранному поставщику')
        return 'inbound'
    if len(rows) >= 10 and text(rows[9][11]).lower() == 'сезонность':
        return 'seasonality'
    if any(month_date(v) for v in rows[0]):
        if any('остаток' in text(v).lower() for row in rows[1:3] for v in row) or 'остат' in filename.lower():
            return 'monthly_stock'
        return 'monthly_sales'
    raise ImportFailure('Неизвестная структура XLSX: ' + filename)


def import_known_workbooks(files, supplier):
    if supplier not in ('systeme', 'iek') or not files:
        raise ImportFailure('Выберите Systeme Electric или IEK и хотя бы одну книгу')
    issues, products, units = [], {}, defaultdict(set)
    transactions, monthly, stocks, inbound, constraints, seasonality, stockouts = [], [], [], [], [], [], []
    roles, books, source_refs = set(), [], []
    if len({PurePosixPath(f.filename.replace('\\','/')).name.lower() for f in files})!=len(files):
        raise ImportFailure('Имена загружаемых книг должны быть уникальны')
    digest = hashlib.sha256(('known-xlsx-v2/' + supplier).encode())
    for file in sorted(files, key=lambda f: f.filename):
        digest.update(file.filename.encode()); digest.update(hashlib.sha256(file.content).digest())
    version = digest.hexdigest()
    def issue(code, message, sku=None, source=None, blocking=False):
        issues.append(Issue(code=code, message=message, sku=sku, source=source,
                            severity='blocking' if blocking else 'warning'))
    def product(sku, name, article='', measure=None, category=None, source=None):
        sku = text(sku)
        if not sku:
            return None
        if sku not in products:
            products[sku] = Product(sku=sku, article=text(article), name=text(name) or sku,
                                    supplier=supplier, unit=None, category=text(category) or None)
        p = products[sku]
        if article:
            p.article = text(article)
        if category is not None:
            p.category = text(category)
        if source and len(p.sources) < 5 and not any(s.file == source.file for s in p.sources):
            p.sources.append(source)
        if measure:
            units[sku].add(unit(measure))
        return sku
    try:
        for file in sorted(files,key=lambda f:f.filename):
            name = PurePosixPath(file.filename.replace('\\', '/')).name
            try:
                with zipfile.ZipFile(io.BytesIO(file.content)) as archive:
                    if sum(i.file_size for i in archive.infolist()) > 250 * 1024 * 1024:
                        raise ImportFailure('Распакованная книга превышает 250 МиБ')
                wb = load_workbook(io.BytesIO(file.content), read_only=True, data_only=True, keep_links=False)
            except ImportFailure:
                raise
            except Exception as exc:
                raise ImportFailure('Не удалось прочитать XLSX: ' + name) from exc
            books.append(wb)
            visible = [s for s in wb if s.sheet_state == 'visible']
            if len(visible) != 1:
                raise ImportFailure('Ожидается один основной видимый лист: ' + name)
            ws = visible[0]
            headers = list(ws.iter_rows(min_row=1, max_row=12, values_only=True))
            role = identify(headers, supplier, name)
            if role in roles:
                raise ImportFailure('Загружено несколько книг одного типа: ' + role)
            roles.add(role)
            source_refs.append(SourceRef(file=name, sheet=ws.title, note=role))
            if role == 'snapshot':
                # Streaming cells omit comments; this small source must retain them.
                wb = load_workbook(io.BytesIO(file.content), data_only=True, keep_links=False)
                books.append(wb); ws = wb[ws.title]
            def ref(row, col=None, note=None):
                return SourceRef(file=name, sheet=ws.title, row=row,
                                 cell=f'{get_column_letter(col)}{row}' if col else None, note=note)
            if role == 'seasonality':
                for row in ws.iter_rows(min_row=11, max_row=22, values_only=True):
                    month = next((i for i, m in enumerate(MONTHS, 1) if text(row[1]).lower().startswith(m)), None)
                    value = number(row[11])
                    if month and value is not None and value > 0:
                        seasonality.append(SeasonalFactor(supplier=supplier, month=month, factor=value, source=ref(month+10,12)))
                if len(seasonality) != 12:
                    raise ImportFailure('В книге сезонности нужны 12 положительных коэффициентов')
                continue
            seen = set()
            report_date = full_date(name)
            months = [(i, month_date(h)) for i, h in enumerate(headers[0]) if month_date(h)]
            customer_column=next((i for i,h in enumerate(headers[0]) if text(h).lower() in ('customer_id','анонимный id клиента')),None)
            for row_index, cells in enumerate(ws.iter_rows(), 1):
                row = [c.value for c in cells]
                if role == 'stockouts':
                    if row_index==1 or not any(v is not None for v in row):
                        continue
                    start,end=full_date(row[1]),full_date(row[2])
                    if not text(row[0]) or not start or not end or end<start:
                        raise ImportFailure('В интервалах stockout нужны SKU, корректные start и end')
                    stockouts.append(StockoutInterval(sku=text(row[0]),start=start,end=end,source=ref(row_index)))
                    continue
                if role == 'transactions':
                    if row_index == 1 or len(row) < 8:
                        continue
                    when = full_date(row[0])
                    if not when:
                        if text(row[3]):
                            product(row[3],row[4],measure=row[5],source=ref(row_index))
                            issue('INVALID_TRANSACTION_DATE', 'Строка операций без распознанной даты не учтена', text(row[3]), ref(row_index), True)
                        continue
                    sku = product(row[3], row[4], measure=row[5], source=ref(row_index))
                    if not sku:
                        issue('MISSING_SKU', 'Строка операций без кода товара', source=ref(row_index))
                        continue
                    qty = number(row[7])
                    if qty is None:
                        issue('MISSING_QUANTITY', 'Количество операции неизвестно', sku, ref(row_index,8), True)
                    document = text(row[2])
                    doc_type = re.split(r'\s+\d', document, maxsplit=1)[0]
                    doc_id = text(row[1])
                    if not doc_id:
                        issue('MISSING_DOCUMENT_ID', 'Неизвестен номер документа; строка отмечена для проверки', sku, ref(row_index), True)
                        doc_id = 'source-row-' + str(row_index)
                    transactions.append(Transaction(sku=sku, date=when, document_id=doc_id,
                        document_type=doc_type, quantity=qty, unit=unit(row[5]), warehouse=text(row[6]) or None,
                        customer_id=(text(row[customer_column]) or None) if customer_column is not None else None,
                        source=ref(row_index)))
                    continue
                if role.startswith('monthly'):
                    if row_index < (4 if role == 'monthly_stock' else 3):
                        continue
                    if supplier == 'systeme':
                        code, name_col, article_col, unit_col = (2,1,None,3) if role == 'monthly_stock' else (1,0,2,None)
                    else:
                        code, name_col, article_col, unit_col = (2,0,None,1) if role == 'monthly_stock' else (1,0,None,None)
                    if not text(row[code]) or text(row[name_col]).lower().startswith('итого'):
                        continue
                    sku = product(row[code], row[name_col], row[article_col] if article_col is not None else '',
                                  row[unit_col] if unit_col is not None else None, source=ref(row_index))
                    if sku in seen:
                        issue('DUPLICATE_MONTHLY_SKU', 'Повтор товара в месячном отчёте; автоматическое суммирование запрещено', sku, ref(row_index), True)
                        continue
                    seen.add(sku)
                    for col, month in months:
                        monthly.append(MonthlyObservation(sku=sku, month=month, quantity=number(row[col]),
                            kind='opening_stock' if role == 'monthly_stock' else 'sales',
                            is_complete=False, source=ref(row_index,col+1)))
                    continue
                if role == 'constraints':
                    if row_index < (3 if supplier == 'systeme' else 2):
                        continue
                    code, article_col, name_col = (2,3,1) if supplier == 'systeme' else (1,2,3)
                    if not text(row[code]):
                        continue
                    sku = product(row[code],row[name_col],row[article_col],source=ref(row_index))
                    value = number(row[4])
                    if sku in seen:
                        issue('DUPLICATE_CONSTRAINT', 'Повтор правила заказа; требуется проверка', sku, ref(row_index), True)
                        continue
                    seen.add(sku)
                    if value is None or value <= 0:
                        issue('INVALID_CONSTRAINT', 'Минимум или кратность не подтверждены числом', sku, ref(row_index,5), True)
                    else:
                        constraints.append(OrderConstraint(sku=sku, minimum=value if supplier=='iek' else None,
                            multiple=value if supplier=='systeme' else None, unit=None, source=ref(row_index,5)))
                    continue
                if role == 'snapshot':
                    if row_index < 3 or not text(row[2]):
                        continue
                    if report_date is None:
                        raise ImportFailure('Дата снимка должна быть указана в имени книги как ДД.ММ.ГГГГ')
                    sku = product(row[2],row[3],row[1],category=row[4],source=ref(row_index))
                    if sku in seen:
                        issue('DUPLICATE_STOCK', 'Повтор строки снимка остатков', sku, ref(row_index), True)
                        continue
                    seen.add(sku)
                    stocks.append(StockSnapshot(sku=sku,as_of=report_date,available=number(row[51]),
                        on_hand=number(row[49]),reserved=number(row[50]),unit=None,is_current=True,
                        source=ref(row_index,52,'Свободный остаток; резерв повторно не вычитается')))
                    latest=stocks[-1]
                    if all(v is not None for v in (latest.available,latest.on_hand,latest.reserved)) and abs(latest.available-(latest.on_hand-latest.reserved))>1e-6:
                        issue('STOCK_RECONCILIATION','Свободный остаток отличается от остатка минус резерв',sku,ref(row_index,52),True)
                    if row[43] is not None:
                        products[sku].sources.append(ref(row_index,44,'Исходный AR=' + text(row[43]) + '; семантика роста не подтверждена, автоматически не применяется'))
                    qty = number(row[54])
                    if qty is None and row[54] is not None:
                        issue('INVALID_INBOUND', 'Неизвестно количество в пути', sku, ref(row_index,55), True)
                    if qty is not None and qty < 0:
                        issue('INVALID_INBOUND', 'Отрицательная поставка требует проверки', sku, ref(row_index,55), True)
                    if qty is not None and qty > 0:
                        comment = cells[54].comment.text if cells[54].comment else None
                        raw_date = comment if comment else headers[1][54]
                        arrival = incoming_date(raw_date, report_date.year)
                        inbound.append(Inbound(sku=sku,quantity=qty,arrival_date=arrival,unit=None,
                            source=ref(row_index,55,('Комментарий: ' if comment else 'Заголовок: ') + text(raw_date) + '; год из даты отчёта')))
                    continue
                if role == 'inbound':
                    if row_index < 2 or not text(row[0]) or text(row[0]).lower().startswith('итого'):
                        continue
                    sku = product(row[0],row[2],row[1],source=ref(row_index,3,text(row[2])))
                    if sku in seen:
                        issue('DUPLICATE_INBOUND_SKU', 'Повтор товара в книге поставок; требуется сверка', sku, ref(row_index), True)
                    seen.add(sku)
                    if 'бухт' in text(row[2]).lower() or 'метраж' in text(row[2]).lower():
                        issue('PURCHASE_UNIT_UNCONFIRMED', 'Закупка бухтами и учёт метрами: преобразование единиц не подтверждено', sku,ref(row_index,3),True)
                    for col in range(3,min(len(row),9)):
                        qty = number(row[col])
                        if row[col] is not None and (qty is None or qty < 0):
                            issue('INVALID_INBOUND', 'Количество поставки требует проверки', sku,ref(row_index,col+1),True)
                        if qty is not None and qty > 0:
                            inbound.append(Inbound(sku=sku,quantity=qty,arrival_date=full_date(headers[0][col]),
                                unit=None,source=ref(row_index,col+1,text(headers[0][col]))))
        if not products:
            raise ImportFailure('В загруженных книгах нет распознанных товаров')
        if any(interval.sku not in products for interval in stockouts):
            raise ImportFailure('Stockout содержит SKU, которого нет в загруженных товарах')
        dated = [s.as_of for s in stocks] + [t.date for t in transactions]
        cutoff = max(dated) if dated else None
        for observation in monthly:
            observation.is_complete = bool(cutoff and observation.month < cutoff.replace(day=1))
        for sku, p in products.items():
            measures = units[sku]
            p.unit = next(iter(measures)) if len(measures)==1 else None
            if len(measures)>1:
                issue('UNIT_CONFLICT','В источниках разные единицы товара; пересчёт не подтверждён',sku,blocking=True)
            if not p.article:
                issue('MISSING_ARTICLE','Артикул поставщика отсутствует; код товара сохранён',sku)
        for record in stocks + inbound + constraints:
            record.unit = products[record.sku].unit
        if supplier=='iek':
            latest = {}
            for observation in monthly:
                if observation.kind=='opening_stock' and (observation.sku not in latest or observation.month>latest[observation.sku].month):
                    latest[observation.sku]=observation
            for sku, observation in latest.items():
                stocks.append(StockSnapshot(sku=sku,as_of=observation.month,available=observation.quantity,
                    unit=products[sku].unit,is_current=False,source=observation.source))
        coverage = None
        if transactions:
            end=max(t.date for t in transactions)
            start=date(2025,1,1) if end.year>=2025 else min(t.date for t in transactions)
            source=next(s for s in source_refs if s.note=='transactions')
            coverage=TransactionCoverage(start=start,end=end,is_complete=False,source=source)
            issue('COVERAGE_ASSUMPTION','Период журнала принят с 01.01.2025 до последней операции по известной выгрузке; полнота не подтверждена. Дни без операций внутри периода трактуются как отсутствие продаж.',source=source)
        if not seasonality:
            issue('MISSING_SEASONALITY','Книга сезонности не загружена; расчёт использует нейтральный профиль с предупреждением')
        dataset=Dataset(id=supplier+'-'+version[:20],version=version,name=('Systeme Electric' if supplier=='systeme' else 'IEK')+' · импорт Excel',
            suppliers=[supplier],is_synthetic=False,created_at=datetime.now(timezone.utc),sources=source_refs,issues=issues)
        return NormalizedDataset(dataset=dataset,products=sorted(products.values(),key=lambda p:p.sku),
            transactions=transactions,monthly=monthly,stocks=stocks,inbound=inbound,constraints=constraints,
            seasonality=seasonality,stockouts=stockouts,transaction_coverage=coverage)
    except ImportFailure:
        raise
    except Exception as exc:
        raise ImportFailure('Структура книги повреждена или не соответствует известному формату') from exc
    finally:
        for book in books:
            book.close()
