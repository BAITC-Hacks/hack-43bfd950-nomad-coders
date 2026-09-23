from datetime import date, timedelta
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.comments import Comment

from app.config import Settings
from app.ingestion import ImportFailure, SourceWorkbook, import_workbooks
from app.main import create_app


def book(filename, rows=None, setup=None):
    wb=Workbook();ws=wb.active
    for row in rows or []:
        ws.append(row)
    if setup:
        setup(ws)
    stream=BytesIO();wb.save(stream);wb.close()
    return SourceWorkbook(filename,stream.getvalue())


def source_set():
    transactions=[['Дата','Номер','Документ','Код','Номенклатура','Ед.','Склад','Количество']]
    for i in range(90):
        transactions.append([(date(2026,9,22)-timedelta(days=i)).strftime('%d.%m.%Y'),'D'+str(i),
            'Расходная накладная '+str(i),'001_','Вымышленный товар','шт','Алматы',10])
    def snapshot(ws):
        ws.title='TDSheet'
        for cell,value in {'C2':'Код 1с','D2':'Наименование','AZ2':'Свободный остаток','BC2':'СЭ в пути 24.09',
            'C3':'001_','B3':'000-ARTICLE','D3':'Вымышленный товар','E3':2,'AX3':110,'AY3':10,'AZ3':100,'BC3':40}.items():
            ws[cell]=value
        ws['BC3'].comment=Comment('поставка 01.10','Synthetic test')
    def season(ws):
        ws['L10']='СЕЗОННОСТЬ'
        for row,month in enumerate(['янв','фев','мар','апр','май','июн','июл','авг','сен','окт','ноя','дек'],11):
            ws.cell(row,2,month);ws.cell(row,12,1)
    return [
        book('Динамика 2025-2026.xlsx',transactions),
        book('Остатки.xlsx',[['№','Номенклатура','Номенклатура.Код','Ед.изм','авг. 2026','сент. 2026'],
                            [],[],[1,'Вымышленный товар','001_','шт',0,None]]),
        book('Продажи.xlsx',[['Номенклатура','Номенклатура.Код','Артикул','Кратность','авг. 2026','сент. 2026'],
                            ['','','','','Количество','Количество'],['Вымышленный товар','001_','000-ARTICLE',12,310,None]]),
        book('MOQ.xlsx',[['№','Номенклатура','Номенклатура.Код','Артикул','Кратность'],[],
                         [1,'Вымышленный товар','001_','000-ARTICLE',12]]),
        book('Путь 22.09.2026.xlsx',setup=snapshot),
        book('Сезонность.xlsx',setup=season),
    ]


def test_known_layout_comments_and_missing_not_zero():
    files=source_set();data=import_workbooks(files,'systeme')
    assert len(data.products)==1 and data.products[0].sku=='001_'
    assert data.products[0].article=='000-ARTICLE'
    assert data.products[0].unit=='шт'
    assert data.stocks[0].available==100 and data.stocks[0].reserved==10
    assert data.inbound[0].arrival_date==date(2026,10,1)
    assert 'Комментарий' in data.inbound[0].source.note
    assert data.monthly[0].quantity==0 and data.monthly[1].quantity is None
    assert all(not m.is_complete for m in data.monthly if m.month.month==9)
    assert len(data.seasonality)==12
    assert data.constraints[0].minimum is None and data.constraints[0].multiple==12
    assert not data.transaction_coverage.is_complete
    reverse=import_workbooks(list(reversed(files)),'systeme')
    assert data.dataset.id==reverse.dataset.id
    assert data.products==reverse.products
    assert data.transactions==reverse.transactions
    assert data.monthly==reverse.monthly


def test_duplicate_roles_broken_and_unknown_files():
    files=source_set()
    for bundle in [[files[0],files[0]],[SourceWorkbook('bad.xlsx',b'broken')],[book('other.xlsx',[['unknown']])]]:
        with pytest.raises(ImportFailure):
            import_workbooks(bundle,'systeme')


def test_negative_transactions_not_flipped_and_unknown_quantities_reported():
    rows=[['Дата','Номер','Документ','Код','Номенклатура','Ед.','Склад','Количество'],
          ['22.09.2026','0001','Возврат','001_','Synthetic','шт','Алматы',-20],
          ['22.09.2026','0002','Продажа','001_','Synthetic','шт','Алматы',None]]
    data=import_workbooks([book('tx.xlsx',rows)],'systeme')
    assert [t.quantity for t in data.transactions]==[-20,None]
    assert any(i.code=='MISSING_QUANTITY' and i.severity=='blocking' for i in data.dataset.issues)
    assert data.transactions[0].document_id=='0001'


def test_iek_minimum_not_multiple_and_old_stock():
    files=[
        book('MOQ IEK.xlsx',[['№','Код 1с','Артикул поставщика','Наименование','Мин. разр. к отгр.'],
                             [1,'001_','A','Synthetic',20],[2,'002_','B','Other','#N/A']]),
        book('Остатки.xlsx',[['Номенклатура','Ед.','Номенклатура.Код','сент. 2026'],
                             ['','','','Количество'],['','','','нач. остаток'],['Synthetic','шт','001_',50]]),
        book('Путь.xlsx',[['Код 1с','Артикул ИЭК','Наименование','Поступление до 01.10.2026'],
                         ['001_','A','Synthetic',40],['001_','A','Synthetic',10]]),
    ]
    data=import_workbooks(files,'iek')
    assert data.constraints[0].minimum==20 and data.constraints[0].multiple is None
    assert data.stocks[0].as_of==date(2026,9,1) and not data.stocks[0].is_current
    assert data.inbound[0].arrival_date==date(2026,10,1)
    assert any(i.code=='INVALID_CONSTRAINT' for i in data.dataset.issues)
    assert any(i.code=='DUPLICATE_INBOUND_SKU' for i in data.dataset.issues)


def test_real_api_import_calculate_order_zero_csv(tmp_path):
    settings=Settings(tmp_path/'data.sqlite3',False,tmp_path/'dist')
    client=TestClient(create_app(settings))
    sources=source_set()
    def upload():
        return client.post('/api/datasets',data={'supplier':'systeme'},files=[
            ('files',(s.filename,s.content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')) for s in sources])
    first=upload();assert first.status_code==201,first.text
    again=upload()
    assert again.json()==first.json()
    assert len(client.get('/api/datasets').json())==1
    calculation=client.post('/api/calculations',json={'dataset_id':first.json()['id'],'settings':{'as_of':'2026-09-22'}})
    assert calculation.status_code==201,calculation.text
    result=calculation.json()
    rec=result['recommendations'][0]
    assert rec['quantity'] is not None,rec['issues']
    assert not result['is_synthetic']  # XLSX import is distinct from hard-coded demo fixture.
    order=client.post('/api/orders',json={'calculation_id':result['id'],'supplier':'systeme','recommendation_ids':[rec['id']]})
    assert order.status_code==201,order.text
    order=order.json()
    saved=client.patch('/api/orders/'+order['id'],json={'expected_revision':1,'overrides':[
        {'recommendation_id':rec['id'],'quantity':0,'reason':'Синтетический тест импорта'}]})
    assert saved.status_code==200
    approved=client.post('/api/orders/'+order['id']+'/approve',json={'expected_revision':2,'confirmed':True}).json()
    export=client.get('/api/orders/'+order['id']+'/export?revision='+str(approved['revision']))
    assert export.status_code==200
    assert ';0;Синтетический тест импорта' in export.content.decode('utf-8-sig')


def test_optional_anonymized_customer_and_confirmed_stockout():
    tx=book('tx.xlsx',[['Дата','Номер','Документ','Код','Номенклатура','Ед.','Склад','Количество','customer_id'],
        ['22.09.2026','D1','Продажа','001_','Synthetic','шт','Алматы',10,'anon-1']])
    outages=book('stockout.xlsx',[['sku','start','end'],['001_','2026-09-01','2026-09-03']])
    data=import_workbooks([tx,outages],'systeme')
    assert data.transactions[0].customer_id=='anon-1'
    assert data.stockouts[0].start==date(2026,9,1)
    assert data.stockouts[0].end==date(2026,9,3)


def test_unknown_date_and_quantity_are_reported():
    file=book('tx.xlsx',[['Дата','Номер','Документ','Код','Номенклатура','Ед.','Склад','Количество'],
                         [None,'D1','Продажа','001_','Synthetic','шт','Алматы',None]])
    data=import_workbooks([file],'systeme')
    assert data.products[0].sku=='001_'
    assert any(i.code=='INVALID_TRANSACTION_DATE' and i.severity=='blocking' for i in data.dataset.issues)
